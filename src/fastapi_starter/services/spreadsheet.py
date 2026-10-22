"""Quick look into an uploaded ``.xlsx`` (successor of the old pandas/openpyxl helpers.py)."""

from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import IO

from openpyxl import load_workbook

from ..errors import UnsupportedMediaError
from ..schemas import SheetPreview


def _jsonable(value: object) -> object:
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def preview_workbook(
    source: Path | IO[bytes], *, rows: int = 5, sheet: str | None = None
) -> SheetPreview:
    """Header plus the first ``rows`` data rows of a sheet (the active one by default)."""
    try:
        wb = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedMediaError("Not a readable .xlsx workbook.") from exc
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise UnsupportedMediaError(
                f"Sheet {sheet!r} not found. Available: {', '.join(wb.sheetnames)}"
            )
        ws = wb[sheet] if sheet is not None else wb.active
        assert ws is not None
        it = ws.iter_rows(values_only=True)
        header = [_jsonable(v) for v in next(it, ())]
        body = [[_jsonable(v) for v in row] for _, row in zip(range(rows), it, strict=False)]
        return SheetPreview(
            sheet=ws.title,
            header=header,
            rows=body,
            total_rows=max((ws.max_row or 0) - 1, 0),
            total_columns=ws.max_column or 0,
        )
    finally:
        wb.close()
